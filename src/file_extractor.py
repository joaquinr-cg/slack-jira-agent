"""File download and text extraction for Slack file attachments."""

import io
import logging
import urllib.request
from typing import Optional

logger = logging.getLogger(__name__)

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB


def download_slack_file(url_private: str, bot_token: str) -> bytes:
    """Download a file from Slack using the bot token for auth.

    Uses a custom redirect handler that preserves the Authorization
    header through cross-origin redirects (standard HTTP clients strip it).

    Args:
        url_private: The url_private field from the Slack file object.
        bot_token: The Slack bot token for Bearer auth.

    Returns:
        The raw file bytes.

    Raises:
        urllib.error.URLError: If the download fails.
        ValueError: If the response is HTML instead of file bytes.
    """
    class _AuthRedirectHandler(urllib.request.HTTPRedirectHandler):
        def redirect_request(self, req, fp, code, msg, headers, newurl):
            new_req = super().redirect_request(req, fp, code, msg, headers, newurl)
            if new_req is not None:
                new_req.add_unredirected_header(
                    "Authorization", f"Bearer {bot_token}",
                )
            return new_req

    opener = urllib.request.build_opener(_AuthRedirectHandler)
    req = urllib.request.Request(url_private)
    req.add_header("Authorization", f"Bearer {bot_token}")
    with opener.open(req, timeout=30) as resp:
        data = resp.read()

    if data[:15] == b"<!DOCTYPE html>":
        raise ValueError(
            "Received HTML instead of file bytes. "
            "Bot may be missing 'files:read' OAuth scope."
        )
    return data


def extract_text_from_bytes(file_bytes: bytes, mimetype: str, filename: str) -> str:
    """Extract text content from file bytes based on MIME type.

    Args:
        file_bytes: Raw file content.
        mimetype: The MIME type of the file.
        filename: The original filename (for logging/placeholders).

    Returns:
        Extracted text content, or a placeholder string for unsupported types.
    """
    # PDF
    if mimetype == "application/pdf":
        return _extract_pdf(file_bytes, filename)

    # DOCX
    if mimetype in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    ):
        return _extract_docx(file_bytes, filename)

    # Plain text, markdown, CSV
    if mimetype.startswith("text/") or mimetype in (
        "application/csv",
        "application/json",
        "application/xml",
    ):
        return _extract_text(file_bytes, filename)

    # Images
    if mimetype.startswith("image/"):
        return f"[Image file: {filename} - content not extracted]"

    # Unsupported
    return f"[Unsupported file type: {filename} ({mimetype})]"


def _extract_pdf(file_bytes: bytes, filename: str) -> str:
    """Extract text from PDF bytes using PyPDF2."""
    try:
        from PyPDF2 import PdfReader

        reader = PdfReader(io.BytesIO(file_bytes))
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
        extracted = "\n\n".join(pages)
        if not extracted.strip():
            return f"[PDF file: {filename} - no extractable text (possibly scanned/image-based)]"
        return extracted
    except Exception as e:
        logger.error("Failed to extract PDF text from %s: %s", filename, str(e))
        return f"[PDF file: {filename} - extraction failed: {str(e)}]"


def _extract_docx(file_bytes: bytes, filename: str) -> str:
    """Extract text from DOCX bytes using python-docx."""
    try:
        from docx import Document

        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        extracted = "\n\n".join(paragraphs)
        if not extracted.strip():
            return f"[DOCX file: {filename} - no extractable text]"
        return extracted
    except Exception as e:
        logger.error("Failed to extract DOCX text from %s: %s", filename, str(e))
        return f"[DOCX file: {filename} - extraction failed: {str(e)}]"


def _extract_text(file_bytes: bytes, filename: str) -> str:
    """Decode plain text / markdown / CSV bytes."""
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        try:
            return file_bytes.decode("latin-1")
        except Exception as e:
            logger.error("Failed to decode text from %s: %s", filename, str(e))
            return f"[Text file: {filename} - decoding failed]"


def extract_text_from_slack_file(
    file_info: dict, bot_token: str
) -> Optional[dict]:
    """Download a Slack file and extract its text content.

    This is a synchronous function. Call via ``asyncio.to_thread()``
    from async code to avoid blocking the event loop.

    Args:
        file_info: A file object from the Slack message ``files[]`` array.
        bot_token: The Slack bot token.

    Returns:
        A dict with ``filename``, ``mimetype``, and ``extracted_text``,
        or None if the file should be skipped.
    """
    filename = file_info.get("name", "unknown")
    mimetype = file_info.get("mimetype", "application/octet-stream")
    size = file_info.get("size", 0)
    # Prefer url_private_download (raw bytes) over url_private (may return HTML preview)
    url_private = file_info.get("url_private_download") or file_info.get("url_private")

    if not url_private:
        logger.warning("No download URL for file %s, skipping", filename)
        return None

    if size > MAX_FILE_SIZE:
        logger.warning(
            "File %s is %d bytes (> %d limit), skipping", filename, size, MAX_FILE_SIZE
        )
        return {
            "filename": filename,
            "mimetype": mimetype,
            "extracted_text": f"[File too large: {filename} ({size} bytes, limit {MAX_FILE_SIZE})]",
        }

    try:
        file_bytes = download_slack_file(url_private, bot_token)
        extracted_text = extract_text_from_bytes(file_bytes, mimetype, filename)
        logger.info("Extracted text from file %s (%s): %d chars", filename, mimetype, len(extracted_text))
        return {
            "filename": filename,
            "mimetype": mimetype,
            "extracted_text": extracted_text,
        }
    except Exception as e:
        logger.error("Failed to download/extract file %s: %s", filename, str(e))
        return {
            "filename": filename,
            "mimetype": mimetype,
            "extracted_text": f"[Failed to process file: {filename} - {str(e)}]",
        }
